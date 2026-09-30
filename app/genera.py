"""Compilazione del modello Word con i dati della pratica."""

from __future__ import annotations

import copy
import logging
import os
import re
import shutil
import subprocess
import tempfile
from decimal import Decimal
from pathlib import Path

import docx
from docx.enum.text import WD_COLOR_INDEX
from docx.oxml.ns import qn
from docxtpl import DocxTemplate

from . import testo as tx

log = logging.getLogger(__name__)

# i dati mancanti vengono scritti come ⟦descrizione⟧ e poi evidenziati in giallo
APRI, CHIUDI = "⟦", "⟧"
_RE_MANCA = re.compile(APRI + r"[^" + CHIUDI + r"]*" + CHIUDI)

_PICCOLI_F = {1: "una"}


def manca(etichetta: str) -> str:
    return f"{APRI}{etichetta}{CHIUDI}"


def _v(valore, etichetta: str) -> str:
    if valore is None or (isinstance(valore, str) and not valore.strip()):
        return manca(etichetta)
    return str(valore).strip()


def numero_con_lettere(n: int, femminile: bool = False) -> str:
    parola = _PICCOLI_F.get(n) if femminile else None
    return f"{n} ({parola or tx.numero_in_lettere(n)})"


def formato_cf(cf: str) -> str:
    """'GMNMRA64R11L407C' -> 'GMN MRA 64R11 L407C' (come nella bozza)."""
    cf = (cf or "").replace(" ", "").upper()
    if len(cf) != 16:
        return cf
    return f"{cf[:3]} {cf[3:6]} {cf[6:11]} {cf[11:]}"


def _genere_gruppo(persone: list[dict], ruolo: str) -> str:
    if len(persone) > 1:
        tutte_f = all(p.get("sesso") == "F" and p.get("tipo") != "societa" for p in persone)
        return "che saranno in seguito denominate" if tutte_f else "che saranno in seguito denominati"
    if not persone:
        return "che sarà in seguito denominata"
    p = persone[0]
    if p.get("tipo") == "societa" or p.get("sesso") == "F":
        return "che sarà in seguito denominata"
    return "che sarà in seguito denominato"


def _persona(p: dict, ruolo: str, n: int) -> dict:
    f = p.get("sesso") == "F"
    chi = f"{ruolo} {n}"
    luogo = _v(p.get("luogo_nascita"), f"luogo di nascita {chi}")
    if p.get("prov_nascita"):
        luogo += f" ({p['prov_nascita']})"
    res = _v(p.get("res_comune"), f"comune di residenza {chi}")
    if p.get("res_prov"):
        res += f" ({p['res_prov']})"
    res += " in " + _v(p.get("res_indirizzo"), f"indirizzo di residenza {chi}")
    doc_tipo = p.get("doc_tipo") or "C.I."
    maschile_doc = doc_tipo.lower().startswith(("passaport",))
    ente = p.get("doc_ente") or ""
    if ente and not ente.lower().startswith(("dal ", "dalla ", "dallo ", "da ")):
        ente = ("dal " if not ente.lower().startswith(("questura", "prefettura", "motorizzazione")) else "dalla ") + ente
    nominativo = f"{p.get('cognome', '')} {p.get('nome', '')}".strip().upper()
    return {
        "societa": p.get("tipo") == "societa",
        "soc_denominazione": _v(p.get("soc_denominazione"), f"denominazione società {chi}"),
        "soc_sede": _v(p.get("soc_sede"), f"sede società {chi}"),
        "soc_cf": _v(p.get("soc_cf"), f"C.F./P.IVA società {chi}"),
        "soc_rea": (p.get("soc_rea") or "").strip(),
        "soc_qualita": _v(p.get("soc_qualita"), "qualità del firmatario"),
        "del_qualita": "del",
        "art": "La Sig.ra" if f else "Il Sig.",
        "art_min": "la Sig.ra" if f else "il Sig.",
        "nominativo": nominativo or manca(f"cognome e nome {chi}"),
        "nato": "nata" if f else "nato",
        "luogo_nascita": luogo,
        "data_nascita": _v(tx.formato_data(p.get("data_nascita")) or p.get("data_nascita"), f"data di nascita {chi}"),
        "cf": _v(formato_cf(p.get("codice_fiscale", "")), f"codice fiscale {chi}"),
        "doc_tipo": doc_tipo,
        "doc_numero": _v(p.get("doc_numero"), f"numero documento {chi}"),
        "rilasciato": "rilasciato" if maschile_doc else "rilasciata",
        "doc_ente": ente or manca(f"ente di rilascio documento {chi}"),
        "doc_rilascio": _v(tx.formato_data(p.get("doc_rilascio")) or p.get("doc_rilascio"), f"data rilascio documento {chi}"),
        "doc_scadenza": _v(tx.formato_data(p.get("doc_scadenza")) or p.get("doc_scadenza"), f"scadenza documento {chi}"),
        "residenza": res,
        "stato_civile": _v(re.sub(r"^(di )?stato civile\s+", "", p.get("stato_civile", "").strip(), flags=re.I),
                           f"stato civile {chi}"),
    }


def _testo_versamento(v: dict, singolo: bool) -> str:
    """Descrizione di un versamento (caparra o acconto) per l'art. 3."""
    if v.get("testo_libero", "").strip():
        corpo = v["testo_libero"].strip().rstrip(";.")
    else:
        rif = (v.get("riferimento") or "").strip()
        versato = v.get("stato") == "versato"
        scad = tx.formato_data(v.get("scadenza"))
        quando = "già versati" if versato else (f"da versarsi entro il {scad}" if scad else "da versarsi alla firma del presente contratto")
        mod = v.get("modalita")
        if mod == "assegno_agenzia":
            corpo = (f"{quando} a mezzo assegno non trasferibile{(' n. ' + rif) if rif else ''} intestato alla Parte "
                     "Promittente Venditrice e consegnato all’Agenzia mediatrice con l’incarico di custodirlo fino "
                     "all’accettazione della proposta e di consegnarlo poi alla Parte Promittente Venditrice")
        elif mod == "bonifico_agenzia":
            corpo = (f"{quando} dalla Parte Promissaria Acquirente all’Agenzia mediatrice con il compito di "
                     "accreditarli sul conto corrente della Parte Promittente Venditrice")
        elif mod == "assegno":
            corpo = (f"{quando} a mezzo assegno non trasferibile{(' n. ' + rif) if rif else ''} intestato alla "
                     "Parte Promittente Venditrice")
        elif mod == "bonifico":
            corpo = (f"{quando} a mezzo bonifico bancario{(' tratto su ' + rif) if rif else ''} direttamente "
                     "intestato alla Parte Promittente Venditrice")
        else:
            corpo = f"{quando}"
    if singolo:
        return corpo
    return f"{tx.euro_completo(v.get('importo')) or manca('importo')} {corpo}"


def _elenco_con_e(parti: list[str]) -> str:
    if len(parti) <= 1:
        return "".join(parti)
    return "; ".join(parti[:-1]) + " e quanto a " + parti[-1]


def contesto(dati: dict, n_pagine: int | None = None) -> dict:
    imm = dati["immobile"]
    prov = dati["provenienza"]
    acc = dati["accordi"]
    urb = dati["urbanistica"]
    ape = dati["ape"]
    firma = dati["firma"]

    venditori = [_persona(p, "venditore", i + 1) for i, p in enumerate(dati["venditori"])]
    acquirenti = [_persona(p, "acquirente", i + 1) for i, p in enumerate(dati["acquirenti"])]
    if not venditori:
        venditori = [_persona({}, "venditore", 1)]
    if not acquirenti:
        acquirenti = [_persona({}, "acquirente", 1)]

    # immobile: unita' raggruppate per sezione/foglio/particella
    gruppi: list[dict] = []
    for u in imm.get("unita", []):
        chiave = (u.get("sezione", ""), u.get("foglio", ""), u.get("particella", ""))
        intest = []
        if u.get("sezione"):
            intest.append(f"Sezione {u['sezione']}")
        intest.append(f"Foglio {_v(u.get('foglio'), 'foglio')}")
        intest.append(f"Particella {_v(u.get('particella'), 'particella')}")
        g = next((g for g in gruppi if g["_k"] == chiave), None)
        if g is None:
            g = {"_k": chiave, "intestazione": " - ".join(intest), "unita": []}
            gruppi.append(g)
        g["unita"].append({
            "sub": _v(u.get("sub"), "subalterno"),
            "categoria": _v(u.get("categoria"), "categoria"),
            "classe": _v(u.get("classe"), "classe"),
            "consistenza": _v(u.get("consistenza"), "consistenza"),
            "rendita": _v(tx.formato_euro(u.get("rendita")), "rendita"),
            "ultima": False,
        })
    if not gruppi:
        gruppi = [{"intestazione": manca("dati catastali"), "unita": [
            {"sub": manca("sub"), "categoria": manca("cat."), "classe": manca("classe"),
             "consistenza": manca("consistenza"), "rendita": manca("rendita"), "ultima": True}]}]
    gruppi[-1]["unita"][-1]["ultima"] = True
    n_unita = sum(len(g["unita"]) for g in gruppi)
    try:
        n_plan = int(str(imm.get("n_planimetrie") or "0").strip() or 0)
    except ValueError:
        n_plan = 0
    n_plan = n_plan or n_unita

    # pagamenti
    prezzo = tx.parse_importo(acc.get("prezzo"))
    caparre = [v for v in acc.get("versamenti", []) if v.get("tipo") == "caparra" and tx.parse_importo(v.get("importo"))]
    acconti = [v for v in acc.get("versamenti", []) if v.get("tipo") == "acconto" and tx.parse_importo(v.get("importo"))]
    tot_caparra = sum((tx.parse_importo(v["importo"]) for v in caparre), Decimal(0))
    tot_acconti = sum((tx.parse_importo(v["importo"]) for v in acconti), Decimal(0))
    saldo = (prezzo - tot_caparra - tot_acconti) if prezzo is not None else None

    numero = 2
    n_caparra = ""
    if caparre:
        n_caparra = f"3.{numero}"
        numero += 1
    ctx_acconti = []
    for v in acconti:
        ctx_acconti.append({"n": f"3.{numero}", "importo": tx.euro_completo(v["importo"]),
                            "dettaglio": _testo_versamento(v, singolo=True)})
        numero += 1
    n_saldo = f"3.{numero}"
    numero += 1
    n_mutuo = f"3.{numero}"

    if len(caparre) == 1:
        caparra_dettaglio = _testo_versamento(caparre[0], singolo=True)
    else:
        caparra_dettaglio = "di cui " + _elenco_con_e([_testo_versamento(v, singolo=False) for v in caparre])

    # dichiarazione urbanistica
    dich = (urb.get("dichiarazione") or "").strip().rstrip(".")
    if not dich:
        dich = ("le opere di costruzione dell’immobile sono iniziate in data anteriore al 1° settembre 1967"
                if urb.get("ante_67") else
                "l’immobile è stato edificato in forza di " + manca("estremi licenza/concessione/permesso di costruire"))

    ape_presente = bool((ape.get("classe") or "").strip())

    clausole = []
    for i, c in enumerate([c for c in acc.get("clausole", []) if (c or "").strip()]):
        testo_c = c.strip()
        if not testo_c.endswith((".", ";")):
            testo_c += "."
        clausole.append({"n": f"5.{7 + i}", "testo": testo_c})

    persone_fisiche = [p for p in venditori + acquirenti if not p["societa"]]

    return {
        "venditori": venditori,
        "acquirenti": acquirenti,
        "V": {"che_sara": _genere_gruppo(dati["venditori"], "venditori")},
        "A_": {"che_sara": _genere_gruppo(dati["acquirenti"], "acquirenti")},
        "I": {
            "tipologia": _v(imm.get("tipologia"), "tipologia immobile"),
            "comune": _v(imm.get("comune"), "comune immobile"),
            "prov": _v(imm.get("prov"), "prov."),
            "indirizzo": _v(imm.get("indirizzo"), "indirizzo immobile"),
            "descrizione": _v(imm.get("descrizione"), "descrizione immobile (es. appartamento al piano ...)"),
            "gruppi": gruppi,
            "plurale": n_unita > 1,
            "n_planimetrie": n_plan,
            "planimetrie": numero_con_lettere(n_plan, femminile=True),
        },
        "P": {
            "tipo_atto": _v(prov.get("tipo_atto"), "tipo atto"),
            "notaio": _v(prov.get("notaio"), "notaio"),
            "notaio_sede": (prov.get("notaio_sede") or "").strip(),
            "data": _v(tx.formato_data(prov.get("data")) or prov.get("data"), "data atto"),
            "repertorio": _v(prov.get("repertorio"), "n. repertorio"),
            "raccolta": (prov.get("raccolta") or "").strip(),
            "registrazione": (prov.get("registrazione") or "").strip(),
            "trascrizione": _v(prov.get("trascrizione"), "estremi di trascrizione"),
        },
        "A": {
            "prezzo": tx.euro_completo(prezzo) if prezzo is not None else manca("prezzo"),
            "caparra": tx.euro_completo(tot_caparra) if caparre else "",
            "n_caparra": n_caparra,
            "caparra_dettaglio": caparra_dettaglio,
            "acconti": ctx_acconti,
            "n_saldo": n_saldo,
            "saldo": tx.euro_completo(saldo) if saldo is not None else manca("importo a saldo"),
            "saldo_modalita": _v(acc.get("saldo_modalita"), "modalità di pagamento del saldo"),
            "mutuo": bool(acc.get("mutuo_condizione")),
            "n_mutuo": n_mutuo,
            "mutuo_importo": tx.euro_completo(acc.get("mutuo_importo")) or manca("importo mutuo"),
            "mutuo_entro": _v(tx.formato_data(acc.get("mutuo_entro")), "termine per il mutuo"),
            "data_rogito": _v(tx.formato_data(acc.get("data_rogito")), "data rogito"),
            "notaio_scelta": acc.get("notaio_scelta") or "Parte Promissaria Acquirente",
            "locato": acc.get("stato_occupazione") == "locato",
        },
        "U": {
            "dichiarazione": dich,
            "conforme": bool(urb.get("conformita_catastale", True)),
            "agibile": bool(urb.get("agibilita_presente", True)),
            "agibilita": _v(urb.get("agibilita"), "estremi certificato di agibilità"),
        },
        "E": {
            "presente": ape_presente,
            "classe": _v(ape.get("classe"), "classe energetica"),
            "ipe": _v(ape.get("ipe"), "IPE"),
            "tecnico": _v(ape.get("tecnico"), "tecnico certificatore"),
            "data": _v(tx.formato_data(ape.get("data")) or ape.get("data"), "data APE"),
        },
        "clausole": clausole,
        "stato_civile": persone_fisiche,
        "F": {
            "foro": _v(firma.get("foro"), "foro competente"),
            "luogo": _v(firma.get("luogo"), "luogo di firma"),
            "data": _v(tx.formato_data(firma.get("data")), "data di firma"),
            "pagine": numero_con_lettere(n_pagine) if n_pagine else manca("n. pagine"),
        },
    }


# ------------------------------------------------------------------ rendering

def _evidenzia_mancanti(percorso: Path) -> int:
    """Trasforma i marcatori ⟦...⟧ in testo '[● ...]' evidenziato in giallo."""
    doc = docx.Document(percorso)
    trovati = 0

    def paragrafi():
        yield from doc.paragraphs
        for t in doc.tables:
            for riga in t.rows:
                for cella in riga.cells:
                    yield from cella.paragraphs

    for par in paragrafi():
        for run in list(par.runs):
            if APRI not in run.text:
                continue
            pezzi = re.split("(" + _RE_MANCA.pattern + ")", run.text)
            prec = run._r
            for pezzo in pezzi:
                if not pezzo:
                    continue
                nuovo = copy.deepcopy(run._r)
                for t in nuovo.findall(qn("w:t")):
                    nuovo.remove(t)
                prec.addnext(nuovo)
                prec = nuovo
                from docx.text.run import Run
                r = Run(nuovo, par)
                if _RE_MANCA.fullmatch(pezzo):
                    r.text = "[● " + pezzo[1:-1] + "]"
                    r.font.highlight_color = WD_COLOR_INDEX.YELLOW
                    trovati += 1
                else:
                    r.text = pezzo
            run._r.getparent().remove(run._r)
    doc.save(percorso)
    return trovati


def trova_soffice() -> str | None:
    """LibreOffice: variabile SOFFICE_CMD, PATH o percorso standard di Windows."""
    for c in (os.environ.get("SOFFICE_CMD"), shutil.which("soffice"), shutil.which("libreoffice"),
              r"C:\Program Files\LibreOffice\program\soffice.exe",
              r"C:\Program Files (x86)\LibreOffice\program\soffice.exe"):
        if c and Path(c).is_file():
            return c
    return None


def conta_pagine(percorso: Path) -> int | None:
    """Conta le pagine convertendo in PDF con LibreOffice (se installato)."""
    soffice = trova_soffice()
    if not soffice:
        return None
    import pymupdf
    with tempfile.TemporaryDirectory() as tmp:
        try:
            subprocess.run([soffice, "-env:UserInstallation=" + (Path(tmp) / "profilo").as_uri(), "--headless",
                            "--convert-to", "pdf", "--outdir", tmp, str(percorso)],
                           check=True, capture_output=True, timeout=120)
            pdf = Path(tmp) / (percorso.stem + ".pdf")
            with pymupdf.open(pdf) as d:
                return len(d)
        except Exception as e:  # pragma: no cover
            log.warning("Conteggio pagine non riuscito: %s", e)
            return None


def _render(modello: Path, ctx: dict, destinazione: Path) -> None:
    tpl = DocxTemplate(str(modello))
    tpl.render(ctx, autoescape=True)
    tpl.save(str(destinazione))


def genera(dati: dict, modello: Path, destinazione: Path) -> dict:
    """Compila il modello. Restituisce {'mancanti': n, 'pagine': n}."""
    ctx = contesto(dati)
    _render(modello, ctx, destinazione)
    pagine = conta_pagine(destinazione)
    if pagine:
        ctx = contesto(dati, pagine)
        _render(modello, ctx, destinazione)
        # il numero scritto in lettere puo' cambiare l'impaginazione: ricontrolla una volta
        p2 = conta_pagine(destinazione)
        if p2 and p2 != pagine:
            _render(modello, contesto(dati, p2), destinazione)
            pagine = p2
    mancanti = _evidenzia_mancanti(destinazione)
    return {"mancanti": mancanti, "pagine": pagine}
