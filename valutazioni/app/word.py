"""Compilazione del modello Word della valutazione.

Nel modello i dati variabili sono scritti tra parentesi quadre, es. [CLIENTE].
- Un segnaposto qualsiasi viene cercato tra le etichette dell'Excel (es. [ZONA OMI]).
- Alcuni sono calcolati dal programma (vedi `campi_calcolati`).
- I paragrafi con segnaposto "per voce" ([VANO], [PERTINENZA], [PRINCIPIO DI UNICITA],
  [CRITICITA]) vengono ripetuti per ogni voce dell'Excel; più paragrafi consecutivi
  con segnaposto dello stesso gruppo vengono ripetuti insieme.
- [VALORI OMI] e [VALORI COMPARABILI] vengono sostituiti dalle immagini.
Il testo sostituito prende la formattazione del segnaposto. Un segnaposto senza
valore resta visibile ed evidenziato in giallo.
"""

from __future__ import annotations

import copy
import re
from pathlib import Path
from typing import Callable

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import qn
from docx.shared import Emu
from docx.text.paragraph import Paragraph
from PIL import Image

from . import excel as xl
from .excel import DatiExcel, norm

SEGNAPOSTO = re.compile(r"\[([^\[\]]{1,80})\]")

GRUPPI = {
    "vani": {"VANO", "MQ VANO", "PERCENTUALE VANO", "VALORE MQ VANO", "VALORE VANO"},
    "pertinenze": {"PERTINENZA", "MQ PERTINENZA", "QUOTA PERTINENZA", "QUOTA PERTINENZA DI", "VALORE MQ PERTINENZA", "VALORE PERTINENZA"},
    "unicita": {"PRINCIPIO DI UNICITA"},
    "criticita": {"CRITICITA"},
}
IMMAGINI = {"VALORI OMI": "omi", "VALORI COMPARABILI": "comparabili"}
# titoli già presenti nel modello dopo cui inserire le immagini se manca il segnaposto
TITOLI_IMMAGINI = {"omi": ("VALORI OMI",), "comparabili": ("VALORI COMPARABILI", "VALORI DI COMPARAZIONE")}


# ------------------------------------------------------------------ valori

def campi_calcolati(d: DatiExcel, oggi: str) -> dict[str, str | None]:
    def e(v):
        return None if v is None else xl.fmt_euro(v)

    c = d.coefficiente_vetusta
    if c is None:
        indice = None
    elif c <= 1:            # 0,7 = il valore attuale è il 70% del nuovo -> decurtazione del 30%
        indice = xl.fmt_numero((1 - c) * 100, 1)
    else:                   # già scritto come percentuale di decurtazione
        indice = xl.fmt_numero(c, 1)
    data = d.valore("DATA")
    return {
        "DATA": xl.fmt_data(data) if data is not None else oggi,
        "VALORE MQ": e(d.valore_mq),
        "VALORE AL MQ": e(d.valore_mq),
        "VALORE AL MQ GARAGE": e(xl.numero(d.valore("Valore al mq GARAGE"))),
        "TOT MQ TIPOLOGIA": xl.fmt_numero(d.tot_mq),
        "TOT VALORE TIPOLOGIA": e(d.tot_valore_tipologia),
        "VALORE A NUOVO": e(d.valore_nuovo),
        "INDICE DI VETUSTA": indice,
        "VETUSTA": e(d.vetusta),
        "VALORE ATTUALE": e(d.valore_attuale),
        "VALORE COMMERCIALE": e(d.valore_commerciale),
    }


def valore_generico(d: DatiExcel, nome: str) -> str | None:
    v = d.valore(nome)
    if v is None:
        return None
    if isinstance(v, str):
        return v.strip()
    if hasattr(v, "strftime"):
        return xl.fmt_data(v)
    n = xl.numero(v)
    if n is None:
        return str(v)
    return xl.fmt_euro(n) if "VALORE" in nome else xl.fmt_numero(n)


def valori_voce(gruppo: str, voce) -> dict[str, str | None]:
    if gruppo in ("unicita", "criticita"):
        return {"PRINCIPIO DI UNICITA": voce, "CRITICITA": voce}
    mq = voce.mq if isinstance(voce.mq, str) else (None if voce.mq is None else xl.fmt_numero(voce.mq))
    suff = "VANO" if gruppo == "vani" else "PERTINENZA"
    return {
        suff: voce.nome,
        f"MQ {suff}": mq,
        f"PERCENTUALE {suff}": None if voce.quota is None else xl.fmt_numero(voce.quota * 100, 1) + "%",
        f"QUOTA {suff}": xl.fmt_quota(voce.quota),
        # riga del calcolo: "a 2.202 €/mq" per intero, "ad 1/3 di 2.202 €/mq" altrimenti
        f"QUOTA {suff} DI": "a" if xl.fmt_quota(voce.quota) == "per intero" else xl.fmt_quota(voce.quota) + " di",
        f"VALORE MQ {suff}": None if voce.valore_mq is None else xl.fmt_euro(voce.valore_mq),
        f"VALORE {suff}": None if voce.valore is None else xl.fmt_euro(voce.valore),
    }


# ------------------------------------------------------------------ testo nei paragrafi

def _testi(p) -> list:
    return [t for t in p.iter(qn("w:t"))]


def testo_paragrafo(p) -> str:
    return "".join(t.text or "" for t in _testi(p))


def _run_di(t):
    r = t.getparent()
    return r if r.tag == qn("w:r") else None


def _evidenzia(run) -> None:
    rpr = run.find(qn("w:rPr"))
    if rpr is None:
        rpr = run.makeelement(qn("w:rPr"), {})
        run.insert(0, rpr)
    for h in rpr.findall(qn("w:highlight")):
        rpr.remove(h)
    h = rpr.makeelement(qn("w:highlight"), {qn("w:val"): "yellow"})
    # ordine dello schema: highlight va prima di u/effect/vertAlign... basta prima di questi
    dopo = [qn(x) for x in ("w:u", "w:effect", "w:bdr", "w:shd", "w:fitText", "w:vertAlign",
                            "w:rtl", "w:cs", "w:em", "w:lang", "w:eastAsianLayout", "w:specVanish", "w:oMath")]
    for i, figlio in enumerate(rpr):
        if figlio.tag in dopo:
            rpr.insert(i, h)
            break
    else:
        rpr.append(h)


def _imposta_testo(t, testo: str) -> None:
    t.text = testo
    t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")


def sostituisci(p, risolvi: Callable[[str], str | None], mancanti: list[str]) -> None:
    """Sostituisce i segnaposto del paragrafo anche se Word li ha spezzati su più run."""
    saltati: set[int] = set()
    while True:
        testi = _testi(p)
        pieno = "".join(t.text or "" for t in testi)
        trovato = None
        for m in SEGNAPOSTO.finditer(pieno):
            if m.start() not in saltati:
                trovato = m
                break
        if not trovato:
            return
        nome = norm(trovato.group(1))
        valore = risolvi(nome)
        # posizioni dei w:t coinvolti
        pos, inizio, fine = 0, None, None
        for i, t in enumerate(testi):
            n = len(t.text or "")
            if inizio is None and trovato.start() < pos + n:
                inizio = (i, trovato.start() - pos)
            if trovato.end() <= pos + n:
                fine = (i, trovato.end() - pos)
                break
            pos += n
        (ia, oa), (ib, ob) = inizio, fine
        ta, tb = testi[ia], testi[ib]
        prima, dopo = (ta.text or "")[:oa], (tb.text or "")[ob:]
        testo_nuovo = valore if valore is not None else trovato.group(0)
        run_a = _run_di(ta)
        if run_a is None:          # w:t fuori da un run: caso anomalo, sostituzione semplice
            _imposta_testo(ta, prima + testo_nuovo + (dopo if ia == ib else ""))
            for t in testi[ia + 1:ib + 1]:
                _imposta_testo(t, "" if t is not tb else dopo)
        else:
            _imposta_testo(ta, prima)
            for t in testi[ia + 1:ib]:
                _imposta_testo(t, "")
            if ib != ia:
                _imposta_testo(tb, dopo)
            nuovo = copy.deepcopy(run_a)
            for figlio in list(nuovo):
                if figlio.tag != qn("w:rPr"):
                    nuovo.remove(figlio)
            t_nuovo = nuovo.makeelement(qn("w:t"), {})
            _imposta_testo(t_nuovo, testo_nuovo)
            nuovo.append(t_nuovo)
            # il nuovo run va subito dopo il w:t del segnaposto, prima di quanto seguiva nel run
            resto = list(run_a)[list(run_a).index(ta) + 1:]
            run_a.addnext(nuovo)
            if ia == ib and (dopo or resto):
                coda = copy.deepcopy(run_a)
                for figlio in list(coda):
                    if figlio.tag != qn("w:rPr"):
                        coda.remove(figlio)
                if dopo:
                    t_coda = coda.makeelement(qn("w:t"), {})
                    _imposta_testo(t_coda, dopo)
                    coda.append(t_coda)
                for figlio in resto:
                    coda.append(figlio)
                nuovo.addnext(coda)
            if valore is None:
                _evidenzia(nuovo)
        if valore is None:
            mancanti.append(trovato.group(1).strip())
            saltati.add(trovato.start())   # il segnaposto resta nel testo: non va ripreso
        else:
            delta = len(testo_nuovo) - (trovato.end() - trovato.start())
            saltati = {x + delta if x > trovato.start() else x for x in saltati}


# ------------------------------------------------------------------ blocchi ripetuti

def pulisci_copia(p):
    """Toglie da un paragrafo copiato segnalibri e identificativi, che devono essere unici."""
    for tag in ("w:bookmarkStart", "w:bookmarkEnd"):
        for x in list(p.iter(qn(tag))):
            x.getparent().remove(x)
    for att in list(p.attrib):
        if att.endswith("}paraId") or att.endswith("}textId"):
            del p.attrib[att]
    return p


def _gruppo(p) -> str | None:
    nomi = {norm(m.group(1)) for m in SEGNAPOSTO.finditer(testo_paragrafo(p))}
    for g, insieme in GRUPPI.items():
        if nomi & insieme:
            return g
    return None


def ripeti_blocchi(doc, d: DatiExcel, risolvi_generale, mancanti: list[str]) -> None:
    corpo = doc.element.body
    paragrafi = [p for p in corpo.iter(qn("w:p"))]
    i = 0
    while i < len(paragrafi):
        g = _gruppo(paragrafi[i])
        if not g:
            i += 1
            continue
        blocco = [paragrafi[i]]
        while i + len(blocco) < len(paragrafi):
            succ = paragrafi[i + len(blocco)]
            if _gruppo(succ) == g and succ.getprevious() is blocco[-1]:
                blocco.append(succ)
            else:
                break
        voci = getattr(d, g)
        ancora = blocco[0]
        for voce in voci:
            valori = valori_voce(g, voce)

            def risolvi(nome, valori=valori):
                return valori[nome] if nome in valori else risolvi_generale(nome)

            for p in blocco:
                nuovo = pulisci_copia(copy.deepcopy(p))
                ancora.addprevious(nuovo)
                sostituisci(nuovo, risolvi, mancanti)
        for p in blocco:
            p.getparent().remove(p)
        i += len(blocco)


# ------------------------------------------------------------------ immagini

def area_testo(doc) -> tuple[int, int]:
    s = doc.sections[0]
    larg = s.page_width - s.left_margin - s.right_margin
    alt = s.page_height - s.top_margin - s.bottom_margin
    return int(larg), int(alt)


def _paragrafo_immagine(modello_p, parent) -> Paragraph:
    nuovo = pulisci_copia(copy.deepcopy(modello_p))
    for figlio in list(nuovo):
        if figlio.tag != qn("w:pPr"):
            nuovo.remove(figlio)
    ppr = nuovo.find(qn("w:pPr"))
    if ppr is not None:
        for tag in ("w:spacing", "w:ind", "w:numPr"):
            for x in ppr.findall(qn(tag)):
                ppr.remove(x)
    par = Paragraph(nuovo, parent)
    par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    par.paragraph_format.space_before = 0
    par.paragraph_format.space_after = Emu(76200)   # 6 pt
    return par


def inserisci_immagini(doc, p_elem, immagini: list[Path], sostituisci_paragrafo: bool) -> None:
    larg, alt = area_testo(doc)
    alt = int(alt * 0.96)
    parent = doc.paragraphs[0]._parent
    ancora = p_elem
    for img in immagini:
        with Image.open(img) as im:
            w, h = im.size
        scala = min(larg / w, alt / h)
        par = _paragrafo_immagine(p_elem, parent)
        par.add_run().add_picture(str(img), width=Emu(int(w * scala)), height=Emu(int(h * scala)))
        ancora.addnext(par._p)
        ancora = par._p
    if sostituisci_paragrafo:
        p_elem.getparent().remove(p_elem)


def posiziona_immagini(doc, tipo: str, immagini: list[Path], avvisi: list[str]) -> None:
    corpo = doc.element.body
    segnaposto = [k for k, v in IMMAGINI.items() if v == tipo]
    for p in list(corpo.iter(qn("w:p"))):
        if any(norm(m.group(1)) in segnaposto for m in SEGNAPOSTO.finditer(testo_paragrafo(p))):
            if immagini:
                inserisci_immagini(doc, p, immagini, sostituisci_paragrafo=True)
            else:
                for t in _testi(p):
                    r = _run_di(t)
                    if r is not None:
                        _evidenzia(r)
            return
    for p in corpo.iter(qn("w:p")):
        if norm(testo_paragrafo(p)) in TITOLI_IMMAGINI[tipo]:
            if immagini:
                inserisci_immagini(doc, p, immagini, sostituisci_paragrafo=False)
            return
    if immagini:
        avvisi.append(f"Nel modello Word non ho trovato dove inserire le immagini {tipo.upper()} "
                      f"(manca il segnaposto [{segnaposto[0]}]).")


# ------------------------------------------------------------------ carta intestata

def _xml_ancora(inline, cx: int, cy: int, id_: int) -> str:
    graphic = inline.find(qn("a:graphic"))
    from lxml import etree
    g = etree.tostring(graphic, encoding="unicode")
    return (
        '<wp:anchor xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing" '
        'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
        'xmlns:pic="http://schemas.openxmlformats.org/drawingml/2006/picture" '
        'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" '
        'distT="0" distB="0" distL="0" distR="0" simplePos="0" relativeHeight="0" behindDoc="1" '
        'locked="1" layoutInCell="1" allowOverlap="1">'
        '<wp:simplePos x="0" y="0"/>'
        '<wp:positionH relativeFrom="page"><wp:posOffset>0</wp:posOffset></wp:positionH>'
        '<wp:positionV relativeFrom="page"><wp:posOffset>0</wp:posOffset></wp:positionV>'
        f'<wp:extent cx="{cx}" cy="{cy}"/>'
        '<wp:effectExtent l="0" t="0" r="0" b="0"/>'
        '<wp:wrapNone/>'
        f'<wp:docPr id="{id_}" name="Carta intestata"/>'
        '<wp:cNvGraphicFramePr/>'
        f'{g}'
        '</wp:anchor>'
    )


def applica_carta_intestata(doc, immagine: Path) -> None:
    """Mette l'immagine della carta intestata dietro al testo, a tutta pagina, in ogni pagina."""
    impostazioni = doc.settings.element
    pari_dispari = impostazioni.find(qn("w:evenAndOddHeaders")) is not None
    id_ = 9000
    for s in doc.sections:
        intestazioni = [s.header]
        if s.different_first_page_header_footer:
            intestazioni.append(s.first_page_header)
        if pari_dispari:
            intestazioni.append(s.even_page_header)
        for h in intestazioni:
            h.is_linked_to_previous = False
            p = h.paragraphs[0] if h.paragraphs else h.add_paragraph()
            run = p.add_run()
            run.add_picture(str(immagine), width=s.page_width, height=s.page_height)
            drawing = run._r.find(qn("w:drawing"))
            inline = drawing.find(qn("wp:inline"))
            id_ += 1
            ancora = parse_xml(_xml_ancora(inline, int(s.page_width), int(s.page_height), id_))
            drawing.replace(inline, ancora)


# ------------------------------------------------------------------ compilazione

def compila(modello: Path, d: DatiExcel, destinazione: Path, *, oggi: str,
            immagini_omi: list[Path], immagini_comparabili: list[Path],
            carta_intestata: Path | None, extra: dict[str, str] | None = None) -> dict:
    doc = Document(str(modello))
    mancanti: list[str] = []
    avvisi: list[str] = []
    calcolati = campi_calcolati(d, oggi)
    extra = {norm(k): v for k, v in (extra or {}).items()}

    def risolvi(nome: str) -> str | None:
        if nome in IMMAGINI:
            return None
        if nome in extra and extra[nome]:
            return extra[nome]
        if nome in calcolati:
            return calcolati[nome]
        return valore_generico(d, nome)

    ripeti_blocchi(doc, d, risolvi, mancanti)

    posiziona_immagini(doc, "omi", immagini_omi, avvisi)
    posiziona_immagini(doc, "comparabili", immagini_comparabili, avvisi)

    def paragrafi_tutti():
        yield from doc.element.body.iter(qn("w:p"))
        for s in doc.sections:
            for parte in (s.header, s.footer, s.first_page_header, s.first_page_footer):
                if not parte.is_linked_to_previous:
                    yield from parte._element.iter(qn("w:p"))

    for p in list(paragrafi_tutti()):
        if SEGNAPOSTO.search(testo_paragrafo(p)):
            # i segnaposto delle immagini rimasti (senza immagini) non si contano due volte
            sostituisci(p, risolvi, mancanti)

    if carta_intestata:
        applica_carta_intestata(doc, carta_intestata)

    destinazione.parent.mkdir(parents=True, exist_ok=True)
    tmp = destinazione.with_name("~tmp " + destinazione.name)
    doc.save(str(tmp))
    tmp.replace(destinazione)
    mancanti_unici = list(dict.fromkeys(m for m in mancanti if norm(m) not in IMMAGINI))
    return {"mancanti": mancanti_unici, "avvisi": avvisi}
