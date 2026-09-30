"""Tabella Word con i valori OMI letti dal sito dell'Agenzia delle Entrate.

Struttura:
    [logo Agenzia]  Banca dati delle quotazioni immobiliari
                    Risultato interrogazione: 2025 - 2° semestre
    Provincia | TREVISO
    Comune    | TREVISO
    ...
    Destinazione: Residenziale
    Tipologia | Stato conservativo | Valore Mercato (€/mq) Min Max | ... (come sul sito)
    Abitazioni civili | NORMALE | 2100 | 2900 | ...
    Destinazione: Commerciale
    ...
    Fonte: Agenzia delle Entrate - OMI, consultazione del gg/mm/aaaa
Le intestazioni della tabella dei valori (celle unite comprese) sono riprese da quelle del sito.
"""

from __future__ import annotations

import copy
import re
from pathlib import Path

from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Cm, Emu, Pt, RGBColor
from docx.text.paragraph import Paragraph

CARATTERE_PREDEFINITO = "Avenir Next LT Pro Light"
GRIGIO = "EDEDED"
AZZURRO = "DCE6F2"
BORDO = "8C8C8C"
INFO_COMUNI = ("Provincia", "Comune", "Fascia/zona", "Codice di zona", "Microzona catastale n.", "Tipologia prevalente")


def _carattere(p_elem) -> str:
    for rf in p_elem.iter(qn("w:rFonts")):
        nome = rf.get(qn("w:ascii"))
        if nome:
            return nome
    return CARATTERE_PREDEFINITO


def _sfondo(cella, colore: str) -> None:
    tcpr = cella._tc.get_or_add_tcPr()
    for x in tcpr.findall(qn("w:shd")):
        tcpr.remove(x)
    tcpr.append(parse_xml(f'<w:shd {nsdecls("w")} w:val="clear" w:color="auto" w:fill="{colore}"/>'))


def _testo(cella, testo: str, font: str, *, grassetto=False, centro=False, dim=8, corsivo=False):
    for extra in cella.paragraphs[1:]:          # le celle unite portano con sé i paragrafi vuoti
        extra._p.getparent().remove(extra._p)
    p = cella.paragraphs[0]
    for r in list(p.runs):
        r._r.getparent().remove(r._r)
    pf = p.paragraph_format
    pf.space_before = pf.space_after = Pt(1)
    pf.line_spacing = 1.0
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if centro else WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run(testo)
    run.font.name = font
    run._r.get_or_add_rPr().get_or_add_rFonts().set(qn("w:hAnsi"), font)
    run.font.size = Pt(dim)
    run.bold = grassetto
    run.italic = corsivo
    cella.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    return run


def _numero(t: str) -> bool:
    return bool(re.fullmatch(r"[\d.,\s€/-]+|[LN]|[LN]\s*/\s*[LN]", t.strip())) and t.strip() != ""


def _prepara_tabella(tbl, larghezze: list[int]) -> None:
    tblpr = tbl._tbl.tblPr
    for tag in ("w:tblBorders", "w:tblLayout", "w:tblW", "w:jc"):
        for x in tblpr.findall(qn(tag)):
            tblpr.remove(x)
    tblpr.append(parse_xml(f'<w:tblW {nsdecls("w")} w:w="{sum(larghezze) * 1440 // 914400}" w:type="dxa"/>'))
    tblpr.append(parse_xml(f'<w:jc {nsdecls("w")} w:val="center"/>'))
    bordi = "".join(f'<w:{b} w:val="single" w:sz="4" w:space="0" w:color="{BORDO}"/>'
                    for b in ("top", "left", "bottom", "right", "insideH", "insideV"))
    tblpr.append(parse_xml(f'<w:tblBorders {nsdecls("w")}>{bordi}</w:tblBorders>'))
    tblpr.append(parse_xml(f'<w:tblLayout {nsdecls("w")} w:type="fixed"/>'))
    tbl.autofit = False
    griglia = tbl._tbl.tblGrid
    for i, col in enumerate(griglia.findall(qn("w:gridCol"))):
        col.set(qn("w:w"), str(larghezze[i] * 1440 // 914400))


def _larghezze(tabelle: list[dict], colonne: int, totale: int) -> list[int]:
    """Colonne proporzionali alla parola più lunga (intestazioni) e al testo più lungo (valori)."""
    pesi = [7.0] * colonne
    for t in tabelle:
        for x in t["celle"]:
            if x["cs"] != 1:
                continue
            parole = x["t"].split()
            misura = max((len(p) for p in parole), default=0) + 1 if x["r"] < t["intestazione"] else len(x["t"]) + 2
            pesi[x["c"]] = max(pesi[x["c"]], misura)
    somma = sum(pesi)
    out = [int(totale * p / somma) for p in pesi]
    out[-1] += totale - sum(out)
    return out


def _riga_unita(tbl, colonne: int):
    riga = tbl.add_row()
    cella = riga.cells[0]
    if colonne > 1:
        cella = cella.merge(riga.cells[colonne - 1])
    return cella


def inserisci(doc, p_elem, dati: dict, logo: Path | None, data_consultazione: str, larghezza: int) -> None:
    """Sostituisce il paragrafo `p_elem` ([VALORI OMI]) con la tabella e la riga della fonte."""
    font = _carattere(p_elem)
    tabelle = dati["tabelle"]
    colonne = max(t["colonne"] for t in tabelle)
    larghezze = _larghezze(tabelle, colonne, larghezza)
    tbl = doc.add_table(rows=0, cols=colonne)
    _prepara_tabella(tbl, larghezze)

    # ---- testata: logo e titolo
    cella = _riga_unita(tbl, colonne)
    p = cella.paragraphs[0]
    p.paragraph_format.space_before = p.paragraph_format.space_after = Pt(3)
    if logo and logo.exists():
        p.add_run().add_picture(str(logo), height=Cm(1.0))
        p.add_run("   ")
    r = p.add_run("Banca dati delle quotazioni immobiliari")
    r.bold, r.font.size, r.font.name = True, Pt(10), font
    semestre = dati.get("semestre") or next((t["semestre"] for t in tabelle if t.get("semestre")), "")
    if semestre:
        p2 = cella.add_paragraph()
        p2.paragraph_format.space_after = Pt(3)
        r = p2.add_run(f"Risultato interrogazione: {semestre}")
        r.font.size, r.font.name = Pt(8.5), font
    cella.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

    # ---- dati della zona (uguali per tutte le destinazioni)
    info = dict(tabelle[0].get("info", []))
    for chiave in INFO_COMUNI:
        if not info.get(chiave):
            continue
        riga = tbl.add_row()
        _testo(riga.cells[0], chiave, font, grassetto=True)
        _sfondo(riga.cells[0], GRIGIO)
        valore = riga.cells[1].merge(riga.cells[colonne - 1]) if colonne > 2 else riga.cells[min(1, colonne - 1)]
        _testo(valore, info[chiave], font)

    # ---- valori per destinazione
    for t in tabelle:
        dest = t.get("destinazione") or dict(t.get("info", [])).get("Destinazione", "")
        if dest:
            c = _riga_unita(tbl, colonne)
            _testo(c, f"Destinazione: {dest}", font, grassetto=True, dim=9)
            _sfondo(c, AZZURRO)
        base = len(tbl.rows)
        for _ in range(t["righe"]):
            tbl.add_row()
        for x in t["celle"]:
            cs = x["cs"]
            if x["c"] + cs == t["colonne"] and t["colonne"] < colonne:
                cs += colonne - t["colonne"]          # tabella più stretta: l'ultima cella occupa il resto
            cella = tbl.cell(base + x["r"], x["c"])
            if x["rs"] > 1 or cs > 1:
                cella = cella.merge(tbl.cell(base + x["r"] + x["rs"] - 1, x["c"] + cs - 1))
            intest = x["r"] < t["intestazione"] or x["th"]
            _testo(cella, x["t"], font, grassetto=intest, centro=intest or (_numero(x["t"]) and x["c"] > 0),
                   dim=7.5 if intest else 8)
            if intest:
                _sfondo(cella, GRIGIO)

    # ---- righe non spezzate tra due pagine; larghezze su ogni cella (serve a Word)
    for riga in tbl.rows:
        trpr = riga._tr.get_or_add_trPr()
        trpr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))
        for i, cella in enumerate(riga.cells):
            cella.width = Emu(larghezze[min(i, colonne - 1)])

    # ---- al posto del segnaposto: tabella + fonte
    p_elem.addprevious(tbl._tbl)
    fonte = copy.deepcopy(p_elem)
    for figlio in list(fonte):
        if figlio.tag != qn("w:pPr"):
            fonte.remove(figlio)
    par = Paragraph(fonte, doc.paragraphs[0]._parent)
    par.paragraph_format.space_before = Pt(3)
    r = par.add_run("Fonte: Agenzia delle Entrate – Osservatorio del Mercato Immobiliare (OMI)"
                    + (f", consultazione del {data_consultazione}" if data_consultazione else "") + ".")
    r.italic, r.font.size, r.font.name = True, Pt(8), font
    r.font.color.rgb = RGBColor(0x59, 0x59, 0x59)
    p_elem.addprevious(fonte)
    p_elem.getparent().remove(p_elem)
