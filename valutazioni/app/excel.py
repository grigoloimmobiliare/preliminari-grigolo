"""Lettura del foglio Excel della stima (es. "STIMA facile appartamento.xlsx").

Il foglio non ha una struttura rigida: i dati si trovano cercando le etichette
(es. "CLIENTE:", "Valore al mq", "ZONA OMI") e prendendo la prima cella piena alla
loro destra. La tabella delle superfici si riconosce dalle intestazioni
"Mq", "Percentuale", "Valore", "Valore": sopra "TOT VALORE TIPOLOGIA" ci sono i vani
dell'abitazione, sotto le pertinenze.

Le celle con formule vengono lette col valore salvato da Excel; se manca (file mai
aperto in Excel) il valore viene ricalcolato qui con le stesse regole del foglio.
"""

from __future__ import annotations

import re
import unicodedata
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path

import openpyxl


def norm(s) -> str:
    """Etichetta normalizzata: maiuscolo, senza accenti, punteggiatura e spazi doppi."""
    s = unicodedata.normalize("NFKD", str(s)).encode("ascii", "ignore").decode()
    s = re.sub(r"[^A-Za-z0-9/%]+", " ", s.upper())
    return re.sub(r"\s+", " ", s).strip()


# etichette che non vanno mai scambiate per il valore della cella alla loro sinistra
ETICHETTE_NOTE = {norm(x) for x in (
    "ZONA", "TIPOLOGIA", "INDIRIZZO", "CLIENTE", "ZONA OMI", "DATA", "COMUNE", "PROVINCIA",
    "Valore al mq", "Valore al mq GARAGE", "Indice di vetusta", "Mq", "Percentuale", "Valore",
    "TOT VALORE TIPOLOGIA", "TOT MQ TIPOLOGIA", "Valore a Nuovo", "Vetustà", "Valore attuale",
    "Valore commerciale", "Principi di unicità", "Criticità",
)}
ELENCHI = {"PRINCIPI DI UNICITA": "unicita", "CRITICITA": "criticita"}


def numero(v) -> float | None:
    if isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return float(v)
    if isinstance(v, str):
        t = v.strip().replace("€", "").replace("mq", "").replace(" ", "")
        if re.fullmatch(r"-?\d{1,3}(\.\d{3})+(,\d+)?|-?\d+(,\d+)?", t):
            return float(t.replace(".", "").replace(",", "."))
        if re.fullmatch(r"-?\d+\.\d+", t):
            return float(t)
    return None


def testo_voce(s: str) -> str:
    """"CAMERA 1" -> "Camera 1"; le diciture già scritte in minuscolo restano come sono."""
    s = re.sub(r"\s+", " ", str(s)).strip()
    return s[:1].upper() + s[1:].lower() if s.isupper() else s


@dataclass
class Voce:
    nome: str
    mq: float | str | None
    quota: float | None
    valore_mq: float | None
    valore: float | None


@dataclass
class DatiExcel:
    foglio: str
    etichette: dict = field(default_factory=dict)      # etichetta normalizzata -> valore
    originali: dict = field(default_factory=dict)      # etichetta normalizzata -> testo originale
    vani: list[Voce] = field(default_factory=list)
    pertinenze: list[Voce] = field(default_factory=list)
    unicita: list[str] = field(default_factory=list)
    criticita: list[str] = field(default_factory=list)
    avvisi: list[str] = field(default_factory=list)

    def valore(self, etichetta: str):
        v = self.etichette.get(norm(etichetta))
        return None if v is None or (isinstance(v, str) and not v.strip()) else v

    # ------------------------------------------------ valori calcolati

    @property
    def valore_mq(self) -> float | None:
        return numero(self.valore("Valore al mq"))

    @property
    def coefficiente_vetusta(self) -> float | None:
        return numero(self.valore("Indice di vetusta"))

    @property
    def tot_mq(self) -> float:
        v = numero(self.valore("TOT MQ TIPOLOGIA"))
        return v if v is not None else sum(x.mq for x in self.vani if isinstance(x.mq, (int, float)))

    @property
    def tot_valore_tipologia(self) -> float:
        v = numero(self.valore("TOT VALORE TIPOLOGIA"))
        return v if v is not None else sum(x.valore or 0 for x in self.vani)

    @property
    def valore_nuovo(self) -> float:
        v = numero(self.valore("Valore a Nuovo"))
        return v if v is not None else self.tot_valore_tipologia + sum(x.valore or 0 for x in self.pertinenze)

    @property
    def valore_attuale(self) -> float | None:
        v = numero(self.valore("Valore attuale"))
        if v is not None:
            return v
        c = self.coefficiente_vetusta
        return None if c is None else self.valore_nuovo * c

    @property
    def vetusta(self) -> float | None:
        v = numero(self.valore("Vetustà"))
        if v is not None:
            return v
        a = self.valore_attuale
        return None if a is None else self.valore_nuovo - a

    @property
    def valore_commerciale(self) -> float | None:
        v = numero(self.valore("Valore commerciale"))
        if v is not None:
            return v
        a = self.valore_attuale
        return None if a is None else a * 1.10   # "aumentato del 10% rispetto al valore tecnico"


def _foglio_principale(wb):
    """Il foglio con più celle compilate (nel modello: "Foglio2")."""
    return max(wb.worksheets, key=lambda ws: sum(1 for r in ws.iter_rows() for c in r if c.value is not None))


def _cella_etichetta(v) -> bool:
    return isinstance(v, str) and (v.strip().endswith(":") or norm(v) in ETICHETTE_NOTE)


def leggi(percorso: Path) -> DatiExcel:
    wb_v = openpyxl.load_workbook(percorso, data_only=True)
    wb_f = openpyxl.load_workbook(percorso, data_only=False)
    ws = _foglio_principale(wb_v)
    ws_f = wb_f[ws.title]
    dati = DatiExcel(foglio=ws.title)

    def val(r, c):
        """Valore salvato; per una formula senza valore salvato restituisce None."""
        return ws.cell(r, c).value

    formule_senza_valore = 0
    righe = ws.max_row
    colonne = ws.max_column

    # ---- etichette -> prima cella piena a destra
    for r in range(1, righe + 1):
        for c in range(1, colonne + 1):
            v = val(r, c)
            if not isinstance(v, str) or not norm(v):
                continue
            chiave = norm(v)
            if chiave in dati.etichette:
                continue
            valore = None
            for c2 in range(c + 1, colonne + 1):
                v2 = val(r, c2)
                if v2 is None and isinstance(ws_f.cell(r, c2).value, str) and ws_f.cell(r, c2).value.startswith("="):
                    formule_senza_valore += 1
                    break
                if v2 is None or (isinstance(v2, str) and not v2.strip()):
                    continue
                if _cella_etichetta(v2):
                    break
                valore = v2
                break
            dati.etichette[chiave] = valore
            dati.originali[chiave] = v.strip()

    # ---- tabella delle superfici
    intest = None
    for r in range(1, righe + 1):
        riga = [norm(val(r, c) or "") for c in range(1, colonne + 1)]
        if "MQ" in riga and "PERCENTUALE" in riga:
            intest = r
            c_mq = riga.index("MQ") + 1
            c_perc = riga.index("PERCENTUALE") + 1
            c_val = [i + 1 for i, x in enumerate(riga) if x == "VALORE"]
            break
    if intest is None:
        dati.avvisi.append("Tabella delle superfici non trovata (intestazioni Mq / Percentuale / Valore).")
    else:
        c_vmq = c_val[0] if c_val else c_perc + 1
        c_tot = c_val[1] if len(c_val) > 1 else c_vmq + 1
        v_mq = dati.valore_mq
        v_mq_garage = numero(dati.valore("Valore al mq GARAGE"))
        sezione = dati.vani
        for r in range(intest + 1, righe + 1):
            nome = val(r, 1)
            if nome is None or not str(nome).strip():
                if sezione is dati.pertinenze:
                    break
                continue
            chiave = norm(nome)
            if chiave == "TOT VALORE TIPOLOGIA":
                sezione = dati.pertinenze
                continue
            if chiave.startswith("TOT") or chiave in ("VALORE A NUOVO",):
                break
            mq_raw = val(r, c_mq)
            mq = numero(mq_raw)
            quota = numero(val(r, c_perc))
            vmq = numero(val(r, c_vmq))
            if vmq is None:
                garage = any(k in chiave for k in ("GARAGE", "BOX", "AUTO"))
                vmq = v_mq_garage if garage and v_mq_garage is not None else v_mq
            tot = numero(val(r, c_tot))
            if tot is None and mq is not None:
                tot = mq * (quota if quota is not None else 1) * (vmq or 0)
            dati.etichette[chiave] = mq_raw     # [SOGGIORNO] nel Word = mq del soggiorno
            voce = Voce(testo_voce(nome), mq if mq is not None else (str(mq_raw).strip() if mq_raw else None),
                        quota, vmq, tot)
            if (mq or 0) > 0 or (voce.valore or 0) > 0 or (isinstance(voce.mq, str) and (voce.valore or 0) > 0):
                sezione.append(voce)

    # ---- elenchi (principi di unicità, criticità): testi alla destra e sotto l'etichetta
    posizioni = []
    for r in range(1, righe + 1):
        for c in range(1, colonne + 1):
            v = val(r, c)
            if isinstance(v, str) and norm(v) in ELENCHI:
                posizioni.append((r, c, ELENCHI[norm(v)]))
    posizioni.sort()
    for i, (r, c, nome) in enumerate(posizioni):
        fine = posizioni[i + 1][0] - 1 if i + 1 < len(posizioni) else righe
        voci = []
        for r2 in range(r, fine + 1):
            for c2 in range(1, colonne + 1):
                if (r2, c2) == (r, c):
                    continue
                v = val(r2, c2)
                if isinstance(v, (int, float)) and not isinstance(v, bool):
                    v = str(v)
                if isinstance(v, str) and v.strip() and norm(v) not in ETICHETTE_NOTE:
                    voci.append(re.sub(r"^\s*[-•·]\s*", "", v.strip()))
        setattr(dati, nome, voci)

    if formule_senza_valore:
        dati.avvisi.append("Alcune formule non hanno un valore salvato: i totali sono stati ricalcolati dal programma. "
                           "Per sicurezza apri l'Excel, salvalo e rigenera.")
    return dati


# ------------------------------------------------------------ formattazione

def fmt_numero(v: float, decimali: int = 2) -> str:
    s = f"{v:,.{decimali}f}".replace(",", "§").replace(".", ",").replace("§", ".")
    if "," in s:
        s = s.rstrip("0").rstrip(",")
    return s


def fmt_euro(v: float) -> str:
    return fmt_numero(round(v), 0)


def fmt_data(v) -> str:
    if isinstance(v, (datetime, date)):
        return v.strftime("%d/%m/%Y")
    return str(v).strip()


def fmt_quota(q: float | None) -> str:
    """Quota della percentuale di calcolo: 1 -> "per intero", 0,333 -> "ad 1/3"."""
    if q is None or abs(q - 1) < 0.005:
        return "per intero"
    for num, den in ((1, 3), (2, 3), (1, 4), (3, 4)):
        if abs(q - num / den) < 0.005:
            return f"ad {num}/{den}" if num == 1 else f"ai {num}/{den}"
    return f"al {fmt_numero(q * 100, 1)}%"
