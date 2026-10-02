"""Prova del programma già avviato: crea una valutazione, carica un Excel, genera il Word.

Uso: python windows/prova_pacchetto.py http://127.0.0.1:8081 <cartella dati>
"""

import sys
import time
from datetime import datetime
from pathlib import Path
from urllib.parse import unquote

import httpx
import openpyxl

base, dati = sys.argv[1], Path(sys.argv[2])
radice = Path(__file__).resolve().parent.parent

wb = openpyxl.load_workbook(radice / "modelli" / "stima_modello.xlsx")
ws = wb["Foglio2"]
ws["F4"], ws["B4"], ws["G5"], ws["F6"] = "Prova", "Via Roma 1", datetime(2026, 1, 1), "Firma Prova"
for cella, mq in {"B10": 6, "B11": 30}.items():
    ws[cella] = mq
excel = dati.parent / "prova.xlsx"
wb.save(excel)

c = httpx.Client(base_url=base, timeout=60)
r = c.post("/valutazioni", data={"nome": "Prova pacchetto"}, follow_redirects=False)
url = r.headers["location"]
c.post(f"{url}/carica/excel", files={"files": ("prova.xlsx", excel.read_bytes())})
# planimetria: due stanze di misura nota (4 x 3 m e 3 x 3 m)
sys.path.insert(0, str(radice))
sys.path.insert(0, str(radice / "tests"))
from test_planimetria import disegna  # noqa: E402
plan, _ = disegna(dati.parent, ridotta=True)
c.post(f"{url}/carica/planimetria", files={"files": ("planimetria.png", plan.read_bytes(), "image/png")})
c.post(f"{url}/planimetria/calcola", data={"porta_max": "1,1"})
import json  # noqa: E402
stato = json.loads((dati / "valutazioni" / unquote(url.rsplit("/", 1)[1]) / "valutazione.json").read_text("utf-8"))
grandi = [z["n"] for z in stato["planimetria"]["zone"] if z["mq"] > 2]
c.post(f"{url}/planimetria/salva", data={**{f"scegli_{n}": "on" for n in grandi}, "scala": "automatica"})
stato = json.loads((dati / "valutazioni" / unquote(url.rsplit("/", 1)[1]) / "valutazione.json").read_text("utf-8"))
tot = stato["planimetria"]["tot_calpestabile"]
assert 20 < tot < 22, f"superficie calpestabile {tot}"
print("PLANIMETRIA OK:", tot, "mq calpestabili,", stato["planimetria"]["tot_commerciale"], "mq commerciali")

c.post(f"{url}/genera", data={"omi": "no"})
cartella = dati / "valutazioni" / unquote(url.rsplit("/", 1)[1])
for _ in range(60):
    word = list(cartella.glob("Valutazione*.docx"))
    if word:
        break
    time.sleep(1)
assert word, f"Word non creato: {list(cartella.iterdir())}"
assert (dati / "modello" / "valutazione_modello.docx").exists()
assert c.get("/impostazioni").status_code == 200 and c.get("/guida").status_code == 200
print("PROVA OK:", word[0])
