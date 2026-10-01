from app import excel


def test_lettura_modello_compilato(excel_compilato):
    d = excel.leggi(excel_compilato)
    assert d.valore("CLIENTE") == "Sig. Mario Rossi"
    assert d.valore("ZONA OMI") == "B1"
    assert d.valore("TIPOLOGIA") == "Appartamento al piano secondo"
    assert [v.nome for v in d.vani] == ["Ingresso", "Soggiorno", "Cucina", "Camera 1", "Camera 2"]
    assert [v.nome for v in d.pertinenze] == ["Garage", "Terrazzi/Poggioli"]
    assert d.tot_mq == 75.5
    assert d.tot_valore_tipologia == 75.5 * 2202
    assert round(d.valore_nuovo) == round(75.5 * 2202 + 18 * 2202 + 9 * 2202 / 3)
    assert round(d.valore_attuale) == round(d.valore_nuovo * 0.7)
    assert d.unicita == ["Posizione centrale", "Ampia terrazza abitabile"]
    assert d.criticita == ["Assenza di ascensore"]
    assert d.avvisi   # formule senza valore salvato: ricalcolate


def test_modello_vuoto_non_ha_voci():
    from conftest import RADICE
    d = excel.leggi(RADICE / "modelli" / "stima_modello.xlsx")
    assert d.vani == [] and d.pertinenze == []
    assert d.valore("CLIENTE") is None       # cella vuota: non prende l'etichetta vicina


def test_formati():
    assert excel.fmt_euro(166251.4) == "166.251,40"
    assert excel.fmt_euro_mq(4500) == "4.500"
    from datetime import date
    assert excel.fmt_data(date(2026, 8, 24)) == "24 agosto 2026"
    assert excel.fmt_numero(30.5) == "30,5"
    assert excel.fmt_numero(1234) == "1.234"
    assert excel.fmt_quota(1) == "per intero"
    assert excel.fmt_quota(0.333333333) == "ad 1/3"
    assert excel.fmt_quota(0.1) == "al 10%"
    assert excel.numero("2.202,50 €") == 2202.5


def test_a_corpo_vetusta_non_applicata_e_commerciale(excel_compilato):
    """Come nella valutazione reale: indice di vetustà 0, pertinenze a corpo, valore commerciale +5%."""
    import openpyxl
    wb = openpyxl.load_workbook(excel_compilato)
    ws = wb["Foglio2"]
    ws["B7"] = 0
    ws["D19"] = "a corpo"        # Magazzino
    ws["E19"] = 10000
    ws["E22"] = 100000           # Posto auto: solo l'importo
    ws["C32"] = None
    nuovo = 75.5 * 2202 + 18 * 2202 + 9 * 2202 / 3 + 10000 + 100000
    ws["D32"] = round(nuovo * 1.05, 2)
    wb.save(excel_compilato)
    d = excel.leggi(excel_compilato)
    corpo = [v for v in d.pertinenze if v.a_corpo]
    assert [v.nome for v in corpo] == ["Magazzino", "Posto Auto"] and corpo[1].valore == 100000
    assert not d.vetusta_applicata
    assert round(d.valore_attuale, 2) == round(nuovo, 2) and d.vetusta == 0
    assert round(d.aumento_commerciale, 1) == 5.0
