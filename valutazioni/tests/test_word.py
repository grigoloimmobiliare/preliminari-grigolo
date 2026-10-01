from docx import Document
from docx.oxml.ns import qn

from app import excel, immagini, word
from conftest import RADICE, immagine

MODELLO = RADICE / "modelli" / "valutazione_modello.docx"


def _testo(doc):
    return "\n".join(p.text for p in doc.paragraphs)


def dati_omi_prova() -> dict:
    from app import omi
    th = lambda t, rs=1, cs=1: {"t": t, "rs": rs, "cs": cs, "th": True}    # noqa: E731
    td = lambda t: {"t": t, "rs": 1, "cs": 1, "th": False}                  # noqa: E731
    righe = [[th("Tipologia", 2), th("Stato conservativo", 2), th("Valore Mercato (€/mq)", cs=2),
              th("Superficie (L/N)", 2)], [th("Min"), th("Max")],
             [td("Abitazioni civili"), td("NORMALE"), td("2100"), td("2900"), td("L")]]
    info = [("Provincia", "TREVISO"), ("Comune", "TREVISO"), ("Codice di zona", "B1")]
    tab = {"semestre": "2025 - 2° semestre", "info": info, "destinazione": "Residenziale", **omi.griglia(righe)}
    return {"semestre": "2025 - 2° semestre", "tabelle": [tab], "data": "30/09/2026"}


def test_compilazione_completa(tmp_path, excel_compilato, carta_pdf):
    d = excel.leggi(excel_compilato)
    comp = [immagine(tmp_path / "comp.png", 900, 1200)]
    carta = immagini.carta_intestata_png(carta_pdf, tmp_path / "carta.png")
    logo = immagine(tmp_path / "logo.png", 200, 50)
    out = tmp_path / "out.docx"
    r = word.compila(MODELLO, d, out, oggi="01/10/2026", immagini_omi=[], immagini_comparabili=comp,
                     carta_intestata=carta, dati_omi=dati_omi_prova(), logo_omi=logo,
                     extra={"COMUNE": "Treviso", "PROVINCIA": "TV"})
    assert r["mancanti"] == [] and r["avvisi"] == []
    doc = Document(str(out))
    t = _testo(doc)
    assert "[" not in t
    assert "Sig. Mario Rossi" in t and "Treviso (TV), lì 30 settembre 2026" in t
    assert "sito in Via Roma 10 Treviso (TV)." in t
    assert "decurtato di un indice di vetustà del 30%" in t and "non applicare detto indice" not in t
    assert "[SE" not in t and "FINE SE" not in t
    for riga in ("Soggiorno: mq 30,5", "Camera 2: mq 11", "per un totale di mq 75,5",
                 "Garage per un totale di 18 mq calcolati per intero",
                 "Terrazzi/Poggioli per un totale di 9 mq calcolati ad 1/3",
                 "- Posizione centrale", "- Assenza di ascensore", "indice di vetustà del 30%"):
        assert riga in t, riga
    assert "Salotto" not in t and "Tinello" not in t
    # tabella OMI subito dopo il titolo "Valori OMI", dopo le pertinenze e prima dei calcoli
    corpo = list(doc.element.body)
    testi = [word.testo_paragrafo(x) for x in corpo]
    i_titolo = testi.index("Valori OMI")
    assert corpo[i_titolo + 1].tag == qn("w:tbl")
    assert testi[i_titolo + 2].startswith("Fonte: Agenzia delle Entrate")
    assert testi.index("Terrazzi/Poggioli per un totale di 9 mq calcolati ad 1/3 sul valore di mercato;") < i_titolo
    assert i_titolo < testi.index("Per cui andiamo a dare un valore alle varie metrature:")
    tab = doc.tables[0]
    assert tab.rows[0].cells[0]._tc.findall(".//" + qn("w:drawing"))          # logo
    celle = [c.text for r in tab.rows for c in r.cells]
    for v in ("Codice di zona", "B1", "Destinazione: Residenziale", "Valore Mercato (€/mq)", "Min", "2900"):
        assert v in celle, v
    # immagine dei comparabili dopo il suo titolo
    paragrafi = doc.paragraphs
    i_comp = [p.text for p in paragrafi].index("Valori di Comparazione")
    con_img = [i for i, p in enumerate(paragrafi) if p._p.findall(".//" + qn("w:drawing"))]
    assert con_img == [i_comp + 1]
    # carta intestata dietro al testo, a tutta pagina
    ancore = doc.sections[0].header._element.findall(".//" + qn("wp:anchor"))
    assert len(ancore) == 1 and ancore[0].get("behindDoc") == "1"


def test_dati_mancanti_evidenziati(tmp_path):
    from conftest import RADICE as R
    d = excel.leggi(R / "modelli" / "stima_modello.xlsx")    # Excel vuoto
    out = tmp_path / "out.docx"
    r = word.compila(MODELLO, d, out, oggi="01/10/2026", immagini_omi=[], immagini_comparabili=[],
                     carta_intestata=None)
    assert "CLIENTE" in r["mancanti"] and "INDIRIZZO" in r["mancanti"]
    doc = Document(str(out))
    evidenziati = [r.text for p in doc.paragraphs for r in p.runs
                   if r._r.find(qn("w:rPr")) is not None and r._r.find(qn("w:rPr")).find(qn("w:highlight")) is not None]
    assert "[CLIENTE]" in evidenziati and "[VALORI OMI]" in evidenziati
    assert "Treviso (TV), lì 01/10/2026" in _testo(doc)
    assert "Pertinenze:" not in _testo(doc)          # nessuna pertinenza: niente titolo


def test_segnaposto_spezzato_su_piu_run(tmp_path, excel_compilato):
    doc = Document()
    p = doc.add_paragraph()
    for pezzo in ("Cliente: [", "CLIEN", "TE] fine", " e [ZONA OMI]."):
        p.add_run(pezzo).bold = pezzo.startswith("CLIEN")
    doc.save(tmp_path / "m.docx")
    d = excel.leggi(excel_compilato)
    word.compila(tmp_path / "m.docx", d, tmp_path / "o.docx", oggi="", immagini_omi=[],
                 immagini_comparabili=[], carta_intestata=None)
    assert Document(str(tmp_path / "o.docx")).paragraphs[0].text == "Cliente: Sig. Mario Rossi fine e B1."


def test_omi_senza_dati_usa_le_immagini(tmp_path, excel_compilato):
    d = excel.leggi(excel_compilato)
    out = tmp_path / "out.docx"
    word.compila(MODELLO, d, out, oggi="", immagini_omi=[immagine(tmp_path / "o.png")], immagini_comparabili=[],
                 carta_intestata=None)
    doc = Document(str(out))
    paragrafi = doc.paragraphs
    i = [p.text for p in paragrafi].index("Valori OMI")
    assert paragrafi[i + 1]._p.findall(".//" + qn("w:drawing")) and not doc.tables


def test_parti_condizionali_e_a_corpo(tmp_path, excel_compilato):
    import openpyxl
    wb = openpyxl.load_workbook(excel_compilato)
    ws = wb["Foglio2"]
    ws["B7"] = 0
    ws["D19"] = "a corpo"
    ws["E19"] = 10000
    wb.save(excel_compilato)
    d = excel.leggi(excel_compilato)
    out = tmp_path / "o.docx"
    word.compila(MODELLO, d, out, oggi="", immagini_omi=[], immagini_comparabili=[], carta_intestata=None)
    t = _testo(Document(str(out)))
    assert "Magazzino: calcolato a corpo;" in t and "- Magazzino:" in t
    assert "Garage per un totale di 18 mq calcolati per intero" in t
    assert "non applicare detto indice" in t and "Coefficiente di vetustà" not in t
    assert "aumentato del 10%" in t
    assert "[SE" not in t and "FINE SE" not in t


def test_condizioni_annidate():
    from docx import Document as D
    doc = D()
    for x in ("prima", "[SE A]", "a", "[SE NON B]", "non b", "[FINE SE]", "[FINE SE]", "[SE B]", "b", "[FINE SE]"):
        doc.add_paragraph(x)
    el = list(doc.element.body)[:-1]
    tenuti = word.applica_condizioni(el, lambda n: {"A": True, "B": True}[n])
    assert [word.testo_paragrafo(p) for p in tenuti] == ["prima", "a", "b"]
