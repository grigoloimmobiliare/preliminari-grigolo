import sys
from datetime import datetime
from pathlib import Path

import openpyxl
import pytest

RADICE = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(RADICE))
sys.path.insert(0, str(RADICE / "tests"))


@pytest.fixture
def excel_compilato(tmp_path) -> Path:
    """Il modello Excel dell'agenzia compilato come farebbe un agente (formule senza valore salvato)."""
    wb = openpyxl.load_workbook(RADICE / "modelli" / "stima_modello.xlsx")
    ws = wb["Foglio2"]
    ws["B3"] = "Centro storico"
    ws["F3"] = "Appartamento al piano secondo"
    ws["B4"] = "Via Roma 10"
    ws["F4"] = "Sig. Mario Rossi"
    ws["B5"] = "B1"
    ws["G5"] = datetime(2026, 9, 30)
    ws["F6"] = "Grigolo Mattia"
    for cella, mq in {"B10": 6, "B11": 30.5, "B12": 12, "B13": 16, "B14": 11, "B17": 18, "B20": 9}.items():
        ws[cella] = mq
    ws["C35"] = "Posizione centrale"
    ws["C36"] = "Ampia terrazza abitabile"
    ws["C42"] = "Assenza di ascensore"
    p = tmp_path / "stima.xlsx"
    wb.save(p)
    return p


@pytest.fixture
def carta_pdf(tmp_path) -> Path:
    import pymupdf
    d = pymupdf.open()
    pg = d.new_page(width=595.3, height=841.9)
    pg.draw_rect(pymupdf.Rect(0, 0, 595.3, 60), color=None, fill=(0.7, 0.5, 0.3))
    pg.insert_text((40, 40), "CARTA INTESTATA", fontsize=20)
    p = tmp_path / "carta.pdf"
    d.save(p)
    return p


def immagine(p: Path, w=800, h=300, colore=(30, 90, 160)) -> Path:
    from PIL import Image
    Image.new("RGB", (w, h), colore).save(p)
    return p
