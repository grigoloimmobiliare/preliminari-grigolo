from pathlib import Path

import docx
import pytest

from app import genera, modello_dati
from app.main import _form_annidato

MODELLO = Path(__file__).resolve().parent.parent / "modelli" / "preliminare_modello.docx"


def _dati():
    d = modello_dati.dati_vuoti()
    p = modello_dati.persona_vuota()
    p.update(sesso="M", cognome="ROSSI", nome="MARIO", luogo_nascita="Treviso", prov_nascita="TV",
             data_nascita="15/03/1980", codice_fiscale="RSSMRA80C15L407X", doc_numero="CA12345AB",
             doc_ente="Comune di Treviso", doc_rilascio="01/01/2020", doc_scadenza="15/03/2031",
             res_indirizzo="Via Roma n. 1", res_comune="Treviso", res_prov="TV", stato_civile="libero")
    a = dict(p, sesso="F", cognome="BIANCHI", nome="GIULIA", stato_civile="")
    d["venditori"], d["acquirenti"] = [p], [a, dict(a, nome="ANNA")]
    d["immobile"].update(comune="Treviso", prov="TV", indirizzo="Via Verdi n. 2", descrizione="appartamento al primo piano")
    d["immobile"]["unita"] = [
        dict(sezione="D", foglio="1", particella="10", sub="3", categoria="A/2", classe="3", consistenza="5 vani", rendita="400,00"),
        dict(sezione="D", foglio="1", particella="10", sub="8", categoria="C/6", classe="2", consistenza="15 mq", rendita="50,00"),
    ]
    d["provenienza"].update(notaio="Paolo Neri", data="10/03/2024", repertorio="1234")
    d["accordi"].update(prezzo="200.000", data_rogito="31/01/2027")
    d["accordi"]["versamenti"] = [
        dict(modello_dati.versamento_vuoto(), importo="5000", modalita="assegno_agenzia", stato="versato", riferimento="123"),
        dict(modello_dati.versamento_vuoto(), importo="15000", modalita="bonifico", scadenza="01/09/2026"),
        dict(modello_dati.versamento_vuoto("acconto"), importo="10000", modalita="bonifico", scadenza="01/10/2026"),
    ]
    d["accordi"]["clausole"] = ["Clausola di prova"]
    d["firma"]["data"] = "29/09/2026"
    return d


def _testo(percorso):
    return "\n".join(p.text for p in docx.Document(percorso).paragraphs)


@pytest.fixture
def generato(tmp_path, monkeypatch):
    monkeypatch.setattr(genera, "conta_pagine", lambda p: None)
    out = tmp_path / "out.docx"
    ris = genera.genera(_dati(), MODELLO, out)
    return out, ris


def test_contenuto(generato):
    out, ris = generato
    t = _testo(out)
    assert "Il Sig. ROSSI MARIO nato a Treviso (TV) il 15/03/1980" in t
    assert "La Sig.ra BIANCHI GIULIA nata a" in t and "BIANCHI ANNA" in t
    assert "che saranno in seguito denominate \"Parte Promissaria Acquirente\"" in t
    assert "che sarà in seguito denominato \"Parte Promittente Venditrice\"" in t
    assert "Sezione D - Foglio 1 - Particella 10:" in t
    assert "- Sub. 8, Cat. C/6, Classe 2, Consistenza 15 mq, Rendita Euro 50,00." in t
    assert "vengono allegate 2 (due) planimetrie catastali" in t
    assert "Le unità immobiliari sono pervenute" in t
    assert "Euro 200.000,00 (duecentomila/00)" in t
    assert "3.2 quanto a Euro 20.000,00 (ventimila/00) a titolo di caparra" in t
    assert "da versarsi entro il 01/09/2026 a mezzo bonifico" in t
    assert "3.3 quanto a Euro 10.000,00 (diecimila/00) a titolo di acconto prezzo" in t
    assert "3.4 quanto a Euro 170.000,00 (centosettantamila/00), a saldo" in t
    assert "5.7 Clausola di prova." in t
    assert "{{" not in t and "{%" not in t


def test_mancanti_evidenziati(generato):
    out, ris = generato
    doc = docx.Document(out)
    evidenziati = [r.text for p in doc.paragraphs for r in p.runs if r.font.highlight_color]
    assert "[● estremi di trascrizione]" in evidenziati
    assert "[● stato civile acquirente 1]" in evidenziati
    assert ris["mancanti"] == len(evidenziati)


def test_modello_vuoto_non_fallisce(tmp_path, monkeypatch):
    monkeypatch.setattr(genera, "conta_pagine", lambda p: None)
    ris = genera.genera(modello_dati.dati_vuoti(), MODELLO, tmp_path / "vuoto.docx")
    assert ris["mancanti"] > 10


def test_form_annidato():
    d = _form_annidato([("venditori.3.cognome", "A"), ("venditori.1.cognome", "B"), ("x.flag", "0"), ("x.flag", "1"),
                        ("accordi.clausole.7", "c")])
    assert [p["cognome"] for p in d["venditori"]] == ["B", "A"]
    assert d["x"]["flag"] == "1"
    assert d["accordi"]["clausole"] == ["c"]
